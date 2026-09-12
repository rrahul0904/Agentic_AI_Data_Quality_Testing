from __future__ import annotations

from email.message import EmailMessage
import io

from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from agentic_data_platform.knowledge import DocumentIntelligencePipeline, extract_document


def test_html_extracts_visible_text_and_excludes_script_style():
    content = b"""
    <html><head><style>.secret{display:none}</style><script>token='bad'</script></head>
    <body><h1>Incident Runbook</h1><p>Scheduler memory pressure.</p></body></html>
    """
    result = extract_document("runbook.html", content, content_type="text/html")
    assert result.source_type == "html"
    assert "Incident Runbook" in result.text
    assert "Scheduler memory pressure" in result.text
    assert "token='bad'" not in result.text
    assert "display:none" not in result.text


def test_eml_extracts_headers_body_and_attachment_metadata():
    message = EmailMessage()
    message["Subject"] = "Pipeline incident"
    message["From"] = "alerts@example.com"
    message["To"] = "data@example.com"
    message.set_content("Airflow worker was restarted after memory pressure.")
    message.add_attachment(b"binary", maintype="application", subtype="octet-stream", filename="trace.bin")
    result = extract_document("incident.eml", message.as_bytes(), content_type="message/rfc822")
    assert result.source_type == "eml"
    assert "Pipeline incident" in result.text
    assert "Airflow worker was restarted" in result.text
    assert result.metadata["attachment_count"] == 1
    assert result.metadata["attachments"] == ["trace.bin"]
    assert "binary" not in result.text


def _presentation_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(0, 0, 5_000_000, 1_000_000)
    textbox.text_frame.text = "Architecture Overview"
    table = slide.shapes.add_table(2, 2, 0, 1_000_000, 5_000_000, 2_000_000).table
    table.cell(0, 0).text = "Layer"
    table.cell(0, 1).text = "Service"
    table.cell(1, 0).text = "Orchestration"
    table.cell(1, 1).text = "Airflow"

    image_buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (255, 255, 255)).save(image_buffer, format="PNG")
    image_buffer.seek(0)
    slide.shapes.add_picture(image_buffer, 500_000, 3_000_000, width=800_000, height=800_000)

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second_box = second.shapes.add_textbox(0, 0, 5_000_000, 1_000_000)
    second_box.text_frame.text = "Recovery Procedures"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_pptx_extracts_slide_text_tables_page_markers_and_asset_geometry():
    result = extract_document(
        "architecture.pptx",
        _presentation_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert result.source_type == "pptx"
    assert "[Page 1]" in result.text
    assert "[Page 2]" in result.text
    assert "Architecture Overview" in result.text
    assert "Recovery Procedures" in result.text
    assert "Orchestration | Airflow" in result.text
    assert result.metadata["slides"] == 2
    assert result.metadata["tables"] == 1
    assert result.metadata["embedded_images"] == 1
    assert result.metadata["page_semantics"] == "slide number"
    asset = result.metadata["assets"][0]
    assert asset["page"] == 1
    assert asset["bbox_unit"] == "emu"
    assert asset["bbox"]["width"] == 800_000
    assert asset["mime_type"] == "image/png"


def test_document_pipeline_page_selection_is_explicit_and_fails_closed_without_pages():
    pipeline = DocumentIntelligencePipeline()
    selected = pipeline.process("architecture.pptx", _presentation_bytes(), pages=[2])
    assert selected.provenance["page_selection"] == [2]
    assert selected.provenance["page_semantics"] == "slide number"
    assert all(block.page == 2 for block in selected.blocks)
    assert "Recovery Procedures" in selected.chunks[0].text
    assert "Architecture Overview" not in selected.chunks[0].text
    assert selected.assets == ()

    try:
        pipeline.process("plain.md", b"No page layout here", pages=[1])
    except ValueError as exc:
        assert "page provenance" in str(exc)
    else:
        raise AssertionError("page selection must fail when the parser has no page provenance")


def test_xlsx_extracts_sheet_values_without_formula_execution():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quality"
    sheet.append(["rule", "status"])
    sheet.append(["not_null", "PASS"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = extract_document(
        "quality.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert result.source_type == "xlsx"
    assert "[Sheet Quality]" in result.text
    assert "rule | status" in result.text
    assert "not_null | PASS" in result.text
    assert result.metadata["sheet_names"] == ["Quality"]
    assert result.metadata["nonempty_rows"] == 2