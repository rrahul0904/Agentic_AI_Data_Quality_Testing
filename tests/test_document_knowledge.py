from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from agentic_data_platform.knowledge import extract_document


def test_extract_markdown() -> None:
    result = extract_document("rules.md", b"# Revenue\nRefunds must be deducted.")
    assert result.source == "rules.md"
    assert result.source_type == "md"
    assert "Refunds must be deducted" in result.text


def test_extract_docx_paragraphs_and_table() -> None:
    document = Document()
    document.add_paragraph("Reservation business rules")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "status"
    table.rows[0].cells[1].text = "NO_SHOW"
    stream = BytesIO()
    document.save(stream)

    result = extract_document("rules.docx", stream.getvalue())
    assert result.source_type == "docx"
    assert "Reservation business rules" in result.text
    assert "status | NO_SHOW" in result.text


def test_rejects_unknown_binary_type() -> None:
    with pytest.raises(ValueError, match="unsupported document type"):
        extract_document("rules.exe", b"MZ")
