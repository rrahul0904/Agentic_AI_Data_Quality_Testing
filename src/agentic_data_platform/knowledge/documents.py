"""Safe, bounded text extraction for project knowledge documents."""

from __future__ import annotations

import csv
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


_TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".sql",
    ".py",
    ".yml",
    ".yaml",
    ".json",
    ".csv",
}
_SUPPORTED_SUFFIXES = _TEXT_SUFFIXES | {
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".html",
    ".htm",
    ".eml",
}


@dataclass(frozen=True)
class DocumentExtraction:
    source: str
    source_type: str
    text: str
    metadata: dict[str, Any]


def _safe_name(filename: str) -> str:
    value = Path(filename or "document").name.strip()
    if not value or value in {".", ".."}:
        raise ValueError("document filename is required")
    return value


def _decode_text(content: bytes) -> str:
    if b"\x00" in content[:4096]:
        raise ValueError("binary content cannot be indexed as plain text")
    return content.decode("utf-8", errors="replace").strip()


class _VisibleHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() in {"script", "style", "noscript", "svg"}:
            self._hidden_depth += 1
        elif not self._hidden_depth and tag.casefold() in {
            "p",
            "div",
            "br",
            "li",
            "tr",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
        }:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"script", "style", "noscript", "svg"}:
            self._hidden_depth = max(0, self._hidden_depth - 1)
        elif not self._hidden_depth and tag.casefold() in {"p", "div", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._hidden_depth and data.strip():
            self.parts.append(data)


def _html_text(value: str) -> str:
    parser = _VisibleHTMLParser()
    parser.feed(value)
    parser.close()
    lines = [" ".join(line.split()) for line in "".join(parser.parts).splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _pdf_text(content: bytes) -> tuple[str, dict[str, Any]]:
    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    extracted_pages = 0
    image_count = 0
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        try:
            image_count += len(page.images)
        except Exception:
            pass
        if text:
            extracted_pages += 1
            pages.append(f"[Page {index}]\n{text}")
    return "\n\n".join(pages).strip(), {
        "pages": len(reader.pages),
        "pages_with_text": extracted_pages,
        "embedded_images": image_count,
    }


def _docx_text(content: bytes) -> tuple[str, dict[str, Any]]:
    document = Document(io.BytesIO(content))
    blocks: list[str] = []
    paragraph_count = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraph_count += 1
            blocks.append(text)
    table_rows = 0
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_rows += 1
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"[Table {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(blocks).strip(), {
        "paragraphs": paragraph_count,
        "tables": len(document.tables),
        "table_rows": table_rows,
    }


def _pptx_text(content: bytes) -> tuple[str, dict[str, Any]]:
    presentation = Presentation(io.BytesIO(content))
    pages: list[str] = []
    table_count = 0
    image_count = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ).strip()
                if text:
                    blocks.append(text)
            if getattr(shape, "has_table", False):
                table_count += 1
                rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    blocks.append(f"[Table {table_count}]\n" + "\n".join(rows))
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                image_count += 1
        if blocks:
            pages.append(f"[Page {slide_index}]\n" + "\n\n".join(blocks))
    return "\n\n".join(pages).strip(), {
        "slides": len(presentation.slides),
        "tables": table_count,
        "embedded_images": image_count,
        "page_semantics": "slide number",
    }


def _xlsx_text(content: bytes) -> tuple[str, dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: list[str] = []
    nonempty_rows = 0
    nonempty_cells = 0
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and values[-1] == "":
                    values.pop()
                if any(value != "" for value in values):
                    nonempty_rows += 1
                    nonempty_cells += sum(value != "" for value in values)
                    rows.append(" | ".join(values))
            if rows:
                sheets.append(f"[Sheet {worksheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets).strip(), {
        "sheets": len(workbook.sheetnames),
        "sheet_names": list(workbook.sheetnames),
        "nonempty_rows": nonempty_rows,
        "nonempty_cells": nonempty_cells,
    }


def _eml_text(content: bytes) -> tuple[str, dict[str, Any]]:
    message = BytesParser(policy=policy.default).parsebytes(content)
    header_names = ("subject", "from", "to", "cc", "date", "message-id")
    headers = {name: str(message.get(name) or "") for name in header_names}
    blocks = [f"{name.title()}: {value}" for name, value in headers.items() if value]
    attachment_names: list[str] = []
    body_parts = 0
    for part in message.walk():
        disposition = part.get_content_disposition()
        filename = part.get_filename()
        if disposition == "attachment" or filename:
            if filename:
                attachment_names.append(str(filename))
            continue
        content_type = part.get_content_type()
        if content_type not in {"text/plain", "text/html"}:
            continue
        try:
            payload = part.get_content()
        except Exception:
            raw = part.get_payload(decode=True) or b""
            payload = raw.decode(part.get_content_charset() or "utf-8", errors="replace")
        text = str(payload).strip()
        if content_type == "text/html":
            text = _html_text(text)
        if text:
            body_parts += 1
            blocks.append(text)
    return "\n\n".join(blocks).strip(), {
        "headers": headers,
        "body_parts": body_parts,
        "attachments": sorted(set(attachment_names)),
        "attachment_count": len(set(attachment_names)),
    }


def extract_document(
    filename: str,
    content: bytes,
    *,
    content_type: str | None = None,
    max_bytes: int = 20_000_000,
) -> DocumentExtraction:
    """Extract bounded textual context from a user-approved document.

    Modern text/office formats and text-bearing PDFs are supported. Scanned/image-only
    documents intentionally fail closed here and are routed through the explicit OCR/vision
    provider contract by :class:`DocumentIntelligencePipeline`.
    """

    source = _safe_name(filename)
    if not content:
        raise ValueError("document is empty")
    if len(content) > max_bytes:
        raise ValueError(f"document exceeds max_bytes={max_bytes}")

    suffix = Path(source).suffix.casefold()
    if suffix not in _SUPPORTED_SUFFIXES:
        raise ValueError(
            f"unsupported document type {suffix or '<none>'}; "
            f"supported: {', '.join(sorted(_SUPPORTED_SUFFIXES))}"
        )

    metadata: dict[str, Any] = {
        "filename": source,
        "bytes": len(content),
        "content_type": content_type,
        "suffix": suffix,
    }
    try:
        if suffix == ".pdf":
            text, detail = _pdf_text(content)
            source_type = "pdf"
        elif suffix == ".docx":
            text, detail = _docx_text(content)
            source_type = "docx"
        elif suffix == ".pptx":
            text, detail = _pptx_text(content)
            source_type = "pptx"
        elif suffix == ".xlsx":
            text, detail = _xlsx_text(content)
            source_type = "xlsx"
        elif suffix in {".html", ".htm"}:
            text = _html_text(_decode_text(content))
            detail = {}
            source_type = "html"
        elif suffix == ".eml":
            text, detail = _eml_text(content)
            source_type = "eml"
        else:
            text = _decode_text(content)
            detail = {}
            source_type = suffix.lstrip(".") or "text"
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(f"unable to parse {suffix.lstrip('.').upper()} document") from exc

    metadata.update(detail)
    if suffix == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, sort_keys=True)
        except json.JSONDecodeError:
            metadata["parse_warning"] = "invalid_json_indexed_as_text"
    elif suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(text)))
            metadata["rows"] = max(0, len(rows) - 1)
            metadata["columns"] = len(rows[0]) if rows else 0
        except csv.Error:
            metadata["parse_warning"] = "invalid_csv_indexed_as_text"

    if not text.strip():
        if suffix in {".pdf", ".pptx"}:
            raise ValueError(
                f"{suffix.lstrip('.').upper()} contains no extractable text; image-only content "
                "requires the configured OCR/vision provider"
            )
        raise ValueError("document contains no extractable text")

    metadata["characters"] = len(text)
    return DocumentExtraction(
        source=source,
        source_type=source_type,
        text=text,
        metadata=metadata,
    )
